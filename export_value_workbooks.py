#!/usr/bin/env python3

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


COMPETITION_SHEETS = [
    "Junior Half Size Line Follower",
    "Senior Half Size Line Follower",
    "Junior Half Size Wall Follower",
    "Senior Half Size Wall Follower",
    "Junior Wall Follower",
    "Senior Wall Follower",
    "Junior Maze Solver",
    "Senior Maze Solver",
    "Junior Half Size Maze Solver",
    "Senior Half Size Maze Solver",
    "Junior Drag Race",
    "Senior Drag Race",
    "Junior Pursuit Race",
    "Senior Pursuit Race",
]

DATA_EXTRACT_SHEETS = [
    "Drag Race Entrants",
    "RATS DATA",
    "Line Following Entrants",
    "Pursuit Entrants",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create value-only Excel workbooks from the master competition workbook."
    )
    parser.add_argument(
        "--workbook",
        default="UKMARS Contest Registration Processing.xlsx",
        help="Path to the source workbook.",
    )
    parser.add_argument(
        "--competition-output-dir",
        default="competition-excels",
        help="Folder for the combined competition workbook.",
    )
    parser.add_argument(
        "--dataextract-output-dir",
        default="dataextracts",
        help="Folder for the single-sheet data extract workbooks.",
    )
    return parser.parse_args()


def safe_filename(name: str) -> str:
    invalid = '<>:"/\\|?*'
    filename = name
    for char in invalid:
        filename = filename.replace(char, "_")
    return filename


def get_export_column_count(ws) -> int:
    last_col = 0
    col = 1
    while True:
        col_letter = get_column_letter(col)
        hidden = ws.column_dimensions[col_letter].hidden
        if hidden:
            break

        has_value = False
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=col, max_col=col):
            value = row[0].value
            if value is not None and value != "":
                has_value = True
                break

        if not has_value:
            break

        last_col = col
        col += 1

    return last_col


def find_used_bounds(ws, last_col: int) -> tuple[int, int]:
    last_row = 0
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=last_col):
        for cell in row:
            value = cell.value
            if value is None or value == "":
                continue
            if cell.row > last_row:
                last_row = cell.row
    return last_row, last_col


def copy_values(source_ws, target_ws) -> None:
    last_col = get_export_column_count(source_ws)
    last_row, last_col = find_used_bounds(source_ws, last_col)
    if last_row == 0 or last_col == 0:
        return

    for row in source_ws.iter_rows(min_row=1, max_row=last_row, min_col=1, max_col=last_col):
        for cell in row:
            target_ws.cell(row=cell.row, column=cell.column, value=cell.value)

    for col_idx in range(1, last_col + 1):
        col_letter = get_column_letter(col_idx)
        target_ws.column_dimensions[col_letter].width = source_ws.column_dimensions[col_letter].width

    for row_idx in range(1, last_row + 1):
        target_ws.row_dimensions[row_idx].height = source_ws.row_dimensions[row_idx].height


def create_blank_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def has_cached_values(ws, column: int, start_row: int = 2, end_row: int = 200) -> bool:
    for row in range(start_row, min(end_row, ws.max_row) + 1):
        value = ws.cell(row=row, column=column).value
        if value is not None and value != "":
            return True
    return False


def export_combined_workbook(source_wb, output_path: Path) -> None:
    output_wb = create_blank_workbook()
    for sheet_name in COMPETITION_SHEETS:
        if sheet_name not in source_wb.sheetnames:
            raise KeyError(f"Missing sheet in source workbook: {sheet_name}")
        target_ws = output_wb.create_sheet(title=sheet_name)
        copy_values(source_wb[sheet_name], target_ws)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(output_path)


def export_single_sheet_workbooks(source_wb, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    for sheet_name in DATA_EXTRACT_SHEETS:
        if sheet_name not in source_wb.sheetnames:
            raise KeyError(f"Missing sheet in source workbook: {sheet_name}")
        output_wb = create_blank_workbook()
        target_ws = output_wb.create_sheet(title=sheet_name)
        copy_values(source_wb[sheet_name], target_ws)
        output_wb.save(output_dir / f"{safe_filename(sheet_name)}.xlsx")


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.workbook)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Source workbook not found: {workbook_path}")

    source_wb = load_workbook(workbook_path, data_only=True)
    if not has_cached_values(source_wb["RobotLookup Clean"], 2):
        raise RuntimeError(
            "The source workbook does not contain cached formula results. "
            "Open it in Excel, recalculate, save, close it, then run this script again."
        )

    competition_output_dir = Path(args.competition_output_dir)
    dataextract_output_dir = Path(args.dataextract_output_dir)

    export_combined_workbook(
        source_wb,
        competition_output_dir / "Competition Sheets.xlsx",
    )
    export_single_sheet_workbooks(source_wb, dataextract_output_dir)

    print(f"Wrote combined workbook to {competition_output_dir / 'Competition Sheets.xlsx'}")
    print(f"Wrote data extract workbooks to {dataextract_output_dir}")


if __name__ == "__main__":
    main()
