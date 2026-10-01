#!/usr/bin/env python3
"""Build the RATS-input variant while preserving the original workbook."""

import argparse
import csv
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.workbook.properties import CalcProperties


HEADERS = [
    "Entry_ID", "Robot", "C_ID", "Contest", "Challenge", "Class", "Source",
    "Contestant", "Cont_Level", "Team_Information",
]
EVENTS = [
    ("Line Follower", "LF"),
    ("Half Size Line Follower", "HLF"),
    ("Drag Race", "DR"),
    ("Non Contact Wall Follower", "WF"),
    ("Wall Follower", "WF"),
    ("Half Size Wall Follower", "HWF"),
    ("Maze Solver", "MS"),
    ("Half Size Maze Solver", "HMS"),
    ("Pursuit Race", "PUR"),
]
CAPACITY = 1000


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", default="UKMARS Contest Registration Processing.xlsx")
    parser.add_argument("--csv", default="RATS-full entry_extract.csv")
    parser.add_argument("--output", default="UKMARS Contest Registration Processing - RATS.xlsx")
    args = parser.parse_args()
    output = Path(args.output)
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite {output}")
    with open(args.csv, encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        if reader.fieldnames != HEADERS:
            raise ValueError(f"Expected these CSV columns in order: {HEADERS}")
        records = list(reader)
    if len(records) > CAPACITY:
        raise ValueError(f"CSV exceeds the {CAPACITY}-entry workbook capacity")

    wb = load_workbook(args.template)
    wb.properties.creator = "UKMARS Example Data"
    wb.properties.lastModifiedBy = "UKMARS Example Data"
    raw = wb["GoogleFormResponsesValueCopy"]
    raw.title = "RATS Entry Import"
    for row in raw:
        for cell in row:
            cell.value = None
    for col, header in enumerate(HEADERS, 1):
        raw.cell(1, col, header)
    for row, record in enumerate(records, 2):
        for col, header in enumerate(HEADERS, 1):
            raw.cell(row, col, record[header] or None)
    raw.freeze_panes = "A2"
    raw.auto_filter.ref = f"A1:J{len(records) + 1}"
    widths = [14, 28, 12, 30, 30, 16, 20, 28, 16, 35]
    for col, width in zip("ABCDEFGHIJ", widths):
        raw.column_dimensions[col].width = width
    raw["A1"].comment = Comment(
        "Paste the complete RATS CSV including headers at A1. Clear all previous "
        "data in A:J first, especially when the new export has fewer rows. "
        "Capacity: 1000 entries. Recalculate and save in desktop Excel.", "UKMARS")
    raw["F1"].comment = Comment(
        "Source competition class (for example Final). Junior/Senior comes from Cont_Level.",
        "UKMARS")

    for name in ["Mouse Query", "RATS DATA", "Denormalised Responses", "RobotLookup"]:
        del wb[name]
    for name in ["existingRobots", "RobotLookup", "startOfEntries"]:
        if name in wb.defined_names:
            del wb.defined_names[name]

    mapping = wb.create_sheet("Event Mapping")
    mapping.append(["RATS Challenge", "Event Code"])
    for event in EVENTS:
        mapping.append(event)
    mapping.column_dimensions["A"].width = 34
    mapping.column_dimensions["B"].width = 18
    mapping.freeze_panes = "A2"
    mapping["A1"].comment = Comment(
        "Match Challenge names exactly. Supported codes: LF, HLF, DR, WF, HWF, MS, HMS, PUR. "
        "Mapping formulas allow additions through row 100.", "UKMARS")
    for cell in mapping[1]:
        cell.font = Font(bold=True)

    lookup = wb["RobotLookup Clean"]
    headers = {1: "Validation", 3: "Challenge", 5: "Level", 7: "Entry ID",
               36: "Source Row", 37: "Event Code", 38: "Contest",
               39: "C_ID", 40: "Source Class", 41: "Source", 42: "Team Information"}
    for col, header in headers.items():
        lookup.cell(1, col, header)
    for row in range(2, CAPACITY + 2):
        lookup.cell(row, 36,
                    f'=IF(COUNTA(\'RATS Entry Import\'!A{row}:J{row})=0,"",ROW())')
        for dest, source_col in {2:"B", 3:"E", 4:"H", 5:"I", 7:"A", 38:"D",
                                 39:"C", 40:"F", 41:"G", 42:"J"}.items():
            lookup.cell(row, dest,
                        f'=IF($AJ{row}="","",IF(\'RATS Entry Import\'!{source_col}{row}="","",'
                        f"'RATS Entry Import'!{source_col}{row}))")
        lookup.cell(row, 6, f'=IF($AJ{row}="","","")')
        lookup.cell(row, 37,
                    f'=IF($AJ{row}="","",IFERROR(VLOOKUP($C{row},'
                    '\'Event Mapping\'!$A$2:$B$100,2,FALSE),"UNMAPPED"))')
        lookup.cell(row, 1,
                    f'=IF($AJ{row}="","",IF(OR($G{row}="",$B{row}="",$D{row}=""),'
                    f'"MISSING DETAILS",IF(AND($E{row}<>"Junior",$E{row}<>"Senior"),'
                    f'"INVALID LEVEL",IF($AH{row}<>1,"UNMAPPED CHALLENGE","OK"))))')
        for col, code in zip(range(8, 16), ["LF", "HLF", "DR", "WF", "HWF", "MS", "HMS", "PUR"]):
            lookup.cell(row, col, f'=IF($AJ{row}="","",IF($AK{row}="{code}",1,0))')
        # Existing Junior/Senior flags and downstream sort/filter formulas stay intact.
    lookup.freeze_panes = "B2"
    lookup["F1"].comment = Comment("Email is unavailable in this RATS CSV; values are blank.", "UKMARS")
    lookup.column_dimensions["A"].width = 26
    for col in ["AK", "AL", "AM", "AN", "AO", "AP"]:
        lookup.column_dimensions[col].width = 24

    summary = wb["Value Copied Summary"]
    checks = [
        (16, "Imported entry rows", '=COUNT(\'RobotLookup Clean\'!$AJ$2:$AJ$1001)'),
        (17, "Entries requiring attention", '=COUNTIF(\'RobotLookup Clean\'!$A$2:$A$1001,"?*")-COUNTIF(\'RobotLookup Clean\'!$A$2:$A$1001,"OK")'),
        (18, "Nonblank cells beyond capacity", '=COUNTA(\'RATS Entry Import\'!$A$1002:$J$1048576)'),
        (19, "Duplicate entry IDs", '=SUMPRODUCT((\'RATS Entry Import\'!$A$2:$A$1001<>"")*(COUNTIF(\'RATS Entry Import\'!$A$2:$A$1001,\'RATS Entry Import\'!$A$2:$A$1001)>1))'),
    ]
    for row, label, formula in checks:
        summary.cell(row, 2, label)
        summary.cell(row, 4, formula)
        summary.cell(row, 2)._style = copy(summary["B13"]._style)
    summary["B21"] = "Paste CSV into RATS Entry Import; clear old data first."
    summary["B22"] = "Recalculate (Ctrl+Alt+F9), save, and check the validation counts above."
    summary["B23"] = "Email is not supplied. Class=Final is source metadata; Cont_Level selects Junior/Senior."
    summary["B24"] = "Each row represents one contest entry. Event Mapping translates Challenge names."
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True,
                                    forceFullCalc=True, calcOnSave=True)
    wb.active = wb.sheetnames.index("RATS Entry Import")
    wb.save(output)
    print(f"Created {output} with {len(records)} entries. Recalculate and save in Excel.")


if __name__ == "__main__":
    main()
